from pathlib import Path

import openpyxl

from excel_mcp.infrastructure import WorkbookStore, Workspace, sha256
from excel_mcp.server import _apply


def make_book(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount"])
    ws.append(["East", 10])
    ws.append(["West", 20])
    ws.append(["East", 5])
    wb.save(path)


def test_commit_creates_backup_and_reopens(tmp_path: Path):
    path = tmp_path / "book.xlsx"
    make_book(path)
    store = WorkbookStore(Workspace(str(tmp_path)))
    before = sha256(path)
    store.commit(path, before, lambda wb: setattr(wb["Sales"]["B2"], "value", 99))
    assert openpyxl.load_workbook(path, data_only=False)["Sales"]["B2"].value == 99
    assert list(tmp_path.glob("book.xlsx.bak-*"))


def test_hash_mismatch_does_not_replace(tmp_path: Path):
    path = tmp_path / "book.xlsx"
    make_book(path)
    store = WorkbookStore(Workspace(str(tmp_path)))
    try:
        store.commit(path, "wrong", lambda wb: None)
    except ValueError:
        pass
    else:
        raise AssertionError("expected hash mismatch")


def test_commit_creates_missing_sheet(tmp_path: Path):
    path = tmp_path / "book.xlsx"
    make_book(path)
    store = WorkbookStore(Workspace(str(tmp_path)))
    before = sha256(path)

    store.commit(
        path,
        before,
        lambda wb: _apply(
            wb,
            {
                "sheet": "Summary",
                "cell": "A1",
                "range": None,
                "values": [["Region", "Amount"]],
            },
        ),
    )
    workbook = openpyxl.load_workbook(path, data_only=False)
    assert "Summary" in workbook.sheetnames
    assert workbook["Summary"]["A1"].value == "Region"
