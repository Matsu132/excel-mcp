from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import threading
import uuid
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


class SecurityError(ValueError):
    pass


class Workspace:
    def __init__(self, root: str | None = None) -> None:
        self.root = Path(root or os.environ.get("EXCEL_MCP_WORKSPACE", os.getcwd())).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def resolve(self, raw: str) -> Path:
        p = Path(raw)
        lexical = self.root / p if not p.is_absolute() else p
        current = lexical
        while current != self.root and current != current.parent:
            if current.is_symlink():
                raise SecurityError("symlink traversal is denied")
            current = current.parent
        candidate = lexical.resolve(strict=False)
        try:
            candidate.relative_to(self.root)
        except ValueError as e:
            raise SecurityError("path is outside workspace") from e
        if candidate.exists() and candidate.is_symlink():
            raise SecurityError("symlink traversal is denied")
        for parent in candidate.parents:
            if parent == self.root:
                break
            if parent.is_symlink():
                raise SecurityError("symlink traversal is denied")
        if candidate.suffix.lower() != ".xlsx":
            raise SecurityError("only .xlsx is supported")
        return candidate


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


class WorkbookStore:
    def __init__(self, workspace: Workspace) -> None:
        self.workspace = workspace
        self._locks: dict[str, threading.Lock] = {}

    def path(self, raw: str) -> Path:
        return self.workspace.resolve(raw)

    def load(self, path: Path) -> Any:
        if not path.exists():
            raise FileNotFoundError(path)
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        if getattr(wb, "_external_links", []):
            wb.close()
            raise SecurityError("external links are not supported")
        return wb

    def lock(self, path: Path) -> threading.Lock:
        return self._locks.setdefault(str(path), threading.Lock())

    def commit(self, path: Path, expected_hash: str, mutate: Callable[[Any], None]) -> None:
        lock = self.lock(path)
        if not lock.acquire(timeout=10):
            raise TimeoutError("workbook is locked")
        try:
            if sha256(path) != expected_hash:
                raise SecurityError("workbook changed since preview")
            backup = path.with_name(
                f"{path.name}.bak-{datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')}"
            )
            shutil.copy2(path, backup)
            wb = self.load(path)
            mutate(wb)
            fd, temp = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=".xlsx", dir=path.parent)
            os.close(fd)
            try:
                wb.save(temp)
                check = self.load(Path(temp))
                check.close()
                os.replace(temp, path)
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)
        finally:
            lock.release()


class PreviewStore:
    def __init__(self) -> None:
        self.items: dict[str, dict[str, Any]] = {}

    def put(self, value: dict[str, Any]) -> str:
        pid = uuid.uuid4().hex
        self.items[pid] = {**value, "preview_id": pid}
        return pid

    def pop(self, pid: str) -> dict[str, Any]:
        if pid not in self.items:
            raise KeyError("unknown preview_id")
        return self.items.pop(pid)


class Audit:
    def __init__(self, workspace: Workspace) -> None:
        self.file = workspace.root / ".excel-mcp-audit.jsonl"

    def write(self, event: dict[str, Any]) -> None:
        with self.file.open("a", encoding="utf-8") as f:
            f.write(json.dumps(event, default=str) + "\n")


def rows(ws: Any, ref: str | None, max_rows: int) -> list[list[Any]]:
    if ref:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
    else:
        min_col, min_row, max_col, max_row = 1, 1, ws.max_column, min(ws.max_row, max_rows)
    return [
        [ws.cell(r, c).value for c in range(min_col, max_col + 1)]
        for r in range(min_row, min(max_row, min_row + max_rows - 1) + 1)
    ]
